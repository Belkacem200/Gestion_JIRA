"""Export de la vue qualite des tickets en HTML, PPTX et PDF (avec coloration)."""
from __future__ import annotations

import html as _html
from datetime import datetime

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                Paragraph, Spacer)

BLUE_HEX = "#003B7A"
BLUE = RGBColor(0x00, 0x3B, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED_HEX = "#c0392b"
GREEN_HEX = "#27ae60"

# Couleur d'etat (categorie de statut)
STATUS_COLORS = {"Done": "#27ae60", "In Progress": "#e67e22", "To Do": "#8895a7"}


def _hex_rgb(h: str) -> RGBColor:
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _idle_hex(days: int) -> str:
    if days >= 90:
        return RED_HEX
    if days >= 30:
        return "#e67e22"
    return "#556070"


def _people_html(r: dict) -> str:
    """Cellule 'Personnes' pour l'export HTML : responsable, participants, coherence."""
    assignee = r.get("assignee") or ""
    parts = r.get("participants", []) or []
    in_part = r.get("assignee_in_participants")
    resp = _html.escape(assignee) if assignee else '<span class="muted">aucun</span>'
    plist = ", ".join(_html.escape(p) for p in parts) if parts \
        else '<span class="muted">aucun</span>'
    coh = ""
    if in_part is True:
        coh = '<div class="cbad ok">✔ responsable dans participants</div>'
    elif in_part is False:
        coh = '<div class="cbad ko">✘ responsable hors participants</div>'
    return (f'<div class="ppl"><div><b>Resp.</b> {resp}</div>'
            f'<div><b>Part.</b> {plist}</div>{coh}</div>')


def _people_pdf(r: dict) -> str:
    """Contenu 'Personnes' pour l'export PDF (balises reportlab)."""
    assignee = _html.escape(r.get("assignee") or "")
    parts = r.get("participants", []) or []
    in_part = r.get("assignee_in_participants")
    resp = assignee if assignee else '<font color="#999999">aucun</font>'
    plist = ", ".join(_html.escape(p) for p in parts) if parts \
        else '<font color="#999999">aucun</font>'
    coh = ""
    if in_part is True:
        coh = '<br/><font color="#1e7e46">&#10004; resp. dans participants</font>'
    elif in_part is False:
        coh = '<br/><font color="#b02a1e">&#10008; resp. hors participants</font>'
    return f'<b>Resp:</b> {resp}<br/><b>Part:</b> {plist}{coh}'


def _people_text(r: dict) -> str:
    """Contenu 'Personnes' en texte brut multi-lignes (export xlsx/csv)."""
    assignee = r.get("assignee") or "aucun"
    parts = r.get("participants", []) or []
    plist = ", ".join(parts) if parts else "aucun"
    lines = [f"Resp. : {assignee}", f"Part. : {plist}"]
    in_part = r.get("assignee_in_participants")
    if in_part is True:
        lines.append("\u2714 responsable dans participants")
    elif in_part is False:
        lines.append("\u2718 responsable hors participants")
    return "\n".join(lines)


def _tree_text(r: dict) -> str:
    """Arborescence (parents / enfants) en texte pour l'export xlsx."""
    parents = r.get("parents", []) or []
    children = r.get("children", []) or []
    seg = []
    if parents:
        seg.append("\u2191 " + ", ".join(p["key"] for p in parents))
    if children:
        seg.append(f"\u2193 {len(children)} : " + ", ".join(c["key"] for c in children))
    return "\n".join(seg)


# ----------------------------- HTML -----------------------------
def build_html(q: dict, title: str) -> str:
    checks = q["checks"]
    cards = "".join(
        f'<div class="card" style="border-top:3px solid {c["color"]}">'
        f'<div class="n" style="color:{c["color"]}">{c["count"]}</div>'
        f'<div class="l">{c["label"]}</div></div>'
        for c in checks.values()
    )
    rows = ""
    for r in q["issues"]:
        badge_parts = []
        for f in r["flags"]:
            badge_parts.append(
                f'<span class="badge" style="background:{checks[f]["color"]}">'
                f'{checks[f]["label"]}</span>')
            if f == "duplicate" and r.get("duplicates"):
                links = ", ".join(
                    f'<a href="{d["url"]}">{_html.escape(d["key"])}</a>'
                    for d in r["duplicates"])
                badge_parts.append(f'<span class="dupof">= {links}</span>')
        badges = "".join(badge_parts) or '<span class="ok">Conforme ✅</span>'
        fc_color = RED_HEX if r["flag_count"] else GREEN_HEX
        st_color = STATUS_COLORS.get(r["status_category"], "#8895a7")
        idle = r["days_idle"]
        idle_txt = idle if idle >= 0 else "—"
        idle_style = f'color:{_idle_hex(idle)};font-weight:600' if idle >= 30 else ""
        parents = r.get("parents", [])
        children = r.get("children", [])
        if parents or children:
            def _li(a):
                return (f'<li><a href="{a["url"]}">{a["key"]}</a>'
                        f'{" — " + _html.escape(a["summary"]) if a["summary"] else ""}</li>')
            inner = ""
            if parents:
                inner += ('<div><b>↑ Parents</b><ul>'
                          + "".join(_li(p) for p in parents) + '</ul></div>')
            if children:
                inner += ('<div><b>↓ Enfants</b><ul>'
                          + "".join(_li(c) for c in children) + '</ul></div>')
            summ = " ".join(filter(None, [
                f"{len(parents)}↑" if parents else "",
                f"{len(children)}↓" if children else ""]))
            parent_cell = f'<details><summary>{summ}</summary>{inner}</details>'
        else:
            parent_cell = '<span class="muted">—</span>'
        rows += (
            f'<tr>'
            f'<td style="text-align:center"><span class="fc" style="background:{fc_color}">{r["flag_count"]}</span></td>'
            f'<td><a href="{r["url"]}">{r["key"]}</a></td>'
            f'<td>{_html.escape(r["summary"] or "")}</td>'
            f'<td>{_html.escape(r["assignee"] or "—")}</td>'
            f'<td><span class="pill" style="background:{st_color}22;color:{st_color};border:1px solid {st_color}55">{_html.escape(r["status"] or "")}</span></td>'
            f'<td style="text-align:right;{idle_style}">{idle_txt}</td>'
            f'<td>{parent_cell}</td>'
            f'<td>{_people_html(r)}</td>'
            f'<td>{badges}</td></tr>'
        )
    today = datetime.now().strftime("%d/%m/%Y %H:%M")
    return f"""<!DOCTYPE html><html lang="fr"><head><meta charset="UTF-8">
<title>Qualité tickets — {_html.escape(title)}</title>
<style>
 body{{font-family:'Segoe UI',Arial,sans-serif;color:#1a2233;margin:0;background:#f4f7fb;}}
 header{{background:linear-gradient(120deg,#003B7A,#0a5bb5);color:#fff;padding:20px 28px;}}
 header h1{{margin:0;font-size:22px;}} .meta{{opacity:.9;font-size:13px;margin-top:4px;}}
 main{{padding:22px 28px;max-width:1300px;margin:0 auto;}}
 .cards{{display:flex;flex-wrap:wrap;gap:12px;margin-bottom:20px;}}
 .card{{background:#fff;border-radius:10px;padding:12px 16px;min-width:130px;box-shadow:0 1px 4px rgba(0,0,0,.06);}}
 .card .n{{font-size:24px;font-weight:700;}} .card .l{{font-size:12px;color:#667;margin-top:2px;}}
 table{{width:100%;border-collapse:collapse;font-size:12px;background:#fff;border-radius:10px;overflow:hidden;box-shadow:0 1px 4px rgba(0,0,0,.06);}}
 th{{background:#003B7A;color:#fff;text-align:left;padding:8px 10px;}}
 td{{padding:7px 10px;border-top:1px solid #e2e8f0;vertical-align:top;}}
 tr:nth-child(even) td{{background:#f8fafd;}}
 a{{color:#0a5bb5;text-decoration:none;font-weight:600;}}
 .badge{{display:inline-block;padding:2px 8px;border-radius:10px;font-size:10.5px;
   font-weight:600;color:#fff;margin:1px 2px;}}
 .fc{{display:inline-block;min-width:22px;text-align:center;color:#fff;border-radius:11px;padding:1px 7px;font-weight:700;}}
 .pill{{padding:2px 9px;border-radius:10px;font-size:11px;font-weight:600;}}
 .ok{{color:#27ae60;font-weight:600;}} .muted{{color:#999;}}
 details summary{{cursor:pointer;color:#0a5bb5;font-weight:600;}}
 details ul{{margin:6px 0 2px;padding-left:16px;}} details li{{margin:2px 0;}}
 .ppl{{font-size:11px;line-height:1.5;}}
 .cbad{{display:inline-block;margin-top:3px;padding:1px 6px;border-radius:8px;font-size:10px;font-weight:600;}}
 .cbad.ok{{background:#eafaf0;color:#1e7e46;}} .cbad.ko{{background:#fdecea;color:#b02a1e;}}
 .dupof{{font-size:10.5px;color:#555;margin:0 2px;}} .dupof a{{color:#0a5bb5;font-weight:600;}}
</style></head><body>
<header>
 <h1>🧹 Qualité &amp; hygiène des tickets</h1>
 <div class="meta">{_html.escape(title)} — généré le {today} — <b>{q["total"]}</b> tickets analysés,
 <b>{q["with_defects"]}</b> avec défaut(s), {q["clean"]} conformes — seuil inactivité {q["stale_days"]} j</div>
</header>
<main>
<div class="cards">{cards}</div>
<table><thead><tr><th>Déf.</th><th>Ticket</th><th>Sujet</th><th>Assigné</th><th>État</th>
 <th style="text-align:right">Inactif (j)</th><th>Arborescence</th><th>Personnes</th><th>Contrôles échoués</th></tr></thead>
<tbody>{rows}</tbody></table>
</main>
</body></html>"""


# ----------------------------- PPTX -----------------------------
def build_pptx(q: dict, title: str, path: str) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Titre
    s = prs.slides.add_slide(blank)
    _band(s, prs)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(2.4), Inches(12), Inches(2))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = "Qualité & hygiène des tickets"
    r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = BLUE
    p2 = tb.text_frame.add_paragraph()
    r2 = p2.add_run()
    r2.text = (f"{title} — {datetime.now().strftime('%d/%m/%Y')} — "
               f"{q['total']} tickets, {q['with_defects']} avec défaut(s)")
    r2.font.size = Pt(16); r2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Synthese des controles (chiffres colores)
    s = prs.slides.add_slide(blank)
    _title(s, "Répartition des défauts")
    checks = list(q["checks"].values())
    tbl = s.shapes.add_table(len(checks) + 1, 2, Inches(0.6), Inches(1.3),
                             Inches(6.5), Inches(0.35 * (len(checks) + 1))).table
    tbl.cell(0, 0).text = "Contrôle"; tbl.cell(0, 1).text = "Nb tickets"
    for i, c in enumerate(checks, start=1):
        tbl.cell(i, 0).text = c["label"]
        cell = tbl.cell(i, 1)
        cell.text = str(c["count"])
        para = cell.text_frame.paragraphs[0]
        if para.runs:
            para.runs[0].font.bold = True
            para.runs[0].font.color.rgb = _hex_rgb(c["color"])
    _style_table(tbl)

    # Tickets prioritaires (top 30, cellules colorees)
    worst = q["issues"][:30]
    for chunk_start in range(0, len(worst), 15):
        chunk = worst[chunk_start:chunk_start + 15]
        s = prs.slides.add_slide(blank)
        _title(s, "Tickets à corriger en priorité")
        t = s.shapes.add_table(len(chunk) + 1, 4, Inches(0.4), Inches(1.2),
                               Inches(12.5), Inches(0.4 * (len(chunk) + 1))).table
        for c, h in enumerate(["Déf.", "Ticket", "Sujet", "Contrôles échoués"]):
            t.cell(0, c).text = h
        for i, r in enumerate(chunk, start=1):
            # Compteur de defauts (cellule coloree)
            fc = t.cell(i, 0)
            fc.text = str(r["flag_count"])
            fc.fill.solid()
            fc.fill.fore_color.rgb = _hex_rgb(RED_HEX if r["flag_count"] else GREEN_HEX)
            fp = fc.text_frame.paragraphs[0]
            if fp.runs:
                fp.runs[0].font.color.rgb = WHITE
                fp.runs[0].font.bold = True
            # Ticket : lien hypertexte cliquable vers JIRA
            kc = t.cell(i, 1)
            kc.text = r["key"]
            krun = kc.text_frame.paragraphs[0].runs[0]
            if r.get("url"):
                krun.hyperlink.address = r["url"]
            krun.font.color.rgb = _hex_rgb("#0a5bb5")
            krun.font.underline = True
            t.cell(i, 2).text = (r["summary"] or "")[:58]
            # Defauts : chaque libelle avec sa couleur
            fcell = t.cell(i, 3)
            tf = fcell.text_frame
            tf.clear()
            para = tf.paragraphs[0]
            if r["flags"]:
                for j, fcode in enumerate(r["flags"]):
                    run = para.add_run()
                    run.text = ("  •  " if j else "") + q["checks"][fcode]["label"]
                    run.font.size = Pt(8)
                    run.font.bold = True
                    run.font.color.rgb = _hex_rgb(q["checks"][fcode]["color"])
                    # Doublon : ajouter les tickets concernes en liens cliquables
                    if fcode == "duplicate" and r.get("duplicates"):
                        op = para.add_run(); op.text = " (="
                        op.font.size = Pt(8); op.font.color.rgb = _hex_rgb("#555555")
                        for di, d in enumerate(r["duplicates"]):
                            if di:
                                sp = para.add_run(); sp.text = ", "
                                sp.font.size = Pt(8); sp.font.color.rgb = _hex_rgb("#555555")
                            lk = para.add_run(); lk.text = d["key"]
                            lk.font.size = Pt(8); lk.font.underline = True
                            lk.font.color.rgb = _hex_rgb("#0a5bb5")
                            if d.get("url"):
                                lk.hyperlink.address = d["url"]
                        cp = para.add_run(); cp.text = ")"
                        cp.font.size = Pt(8); cp.font.color.rgb = _hex_rgb("#555555")
            else:
                run = para.add_run(); run.text = "Conforme"
                run.font.size = Pt(8); run.font.color.rgb = _hex_rgb(GREEN_HEX)
        _style_table(t, sizes=[1.0, 1.8, 5.5, 4.2], skip_body_color={0, 3})

    prs.save(path)


def _band(slide, prs) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.25))
    shp.fill.solid(); shp.fill.fore_color.rgb = BLUE
    shp.line.fill.background()


def _title(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = BLUE


def _style_table(tbl, sizes=None, skip_body_color=None) -> None:
    skip_body_color = skip_body_color or set()
    for c in range(len(tbl.columns)):
        cell = tbl.cell(0, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
        para = cell.text_frame.paragraphs[0]
        if para.runs:
            para.runs[0].font.color.rgb = WHITE
            para.runs[0].font.bold = True
            para.runs[0].font.size = Pt(11)
    for row in range(1, len(tbl.rows)):
        for c in range(len(tbl.columns)):
            if c in skip_body_color:
                continue
            para = tbl.cell(row, c).text_frame.paragraphs[0]
            for run in para.runs:
                run.font.size = Pt(9)
    if sizes:
        for i, w in enumerate(sizes):
            tbl.columns[i].width = Inches(w)


# ----------------------------- PDF -----------------------------
def build_pdf(q: dict, title: str, path: str) -> None:
    doc = SimpleDocTemplate(path, pagesize=landscape(A4),
                            leftMargin=1 * cm, rightMargin=1 * cm,
                            topMargin=1 * cm, bottomMargin=1 * cm)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], textColor=colors.HexColor(BLUE_HEX),
                        fontSize=18, spaceAfter=6, alignment=0)
    meta = ParagraphStyle("meta", parent=styles["Normal"], fontSize=9,
                          textColor=colors.HexColor("#555555"), spaceAfter=10)
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=7, leading=9)
    cellc = ParagraphStyle("cellc", parent=cell, alignment=1)
    head = ParagraphStyle("head", parent=styles["Normal"], fontSize=8,
                         textColor=colors.white, leading=10)

    elems = [Paragraph("Qualite &amp; hygiene des tickets", h1)]
    elems.append(Paragraph(
        f"{_html.escape(title)} &mdash; {datetime.now().strftime('%d/%m/%Y %H:%M')} "
        f"&mdash; <b>{q['total']}</b> tickets analyses, <b>{q['with_defects']}</b> avec defaut(s), "
        f"{q['clean']} conformes &mdash; seuil inactivite {q['stale_days']} j", meta))

    # Tableau de synthese (compte colore)
    check_rows = [[Paragraph("<b>Controle</b>", head), Paragraph("<b>Nb</b>", head)]]
    for c in q["checks"].values():
        check_rows.append([
            Paragraph(c["label"], cell),
            Paragraph(f'<font color="{c["color"]}"><b>{c["count"]}</b></font>', cellc),
        ])
    summ = Table(check_rows, colWidths=[6 * cm, 2 * cm])
    summ.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BLUE_HEX)),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f6fb")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elems.append(summ)
    elems.append(Spacer(1, 12))

    # Tableau des tickets (defauts colores, compteur colore, inactivite coloree)
    header = [Paragraph("<b>Def.</b>", head), Paragraph("<b>Ticket</b>", head),
              Paragraph("<b>Sujet</b>", head), Paragraph("<b>Assigne</b>", head),
              Paragraph("<b>Etat</b>", head),
              Paragraph("<b>Inactif</b>", head), Paragraph("<b>Personnes</b>", head),
              Paragraph("<b>Controles echoues</b>", head)]
    data = [header]
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BLUE_HEX)),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#dddddd")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    for idx, r in enumerate(q["issues"], start=1):
        flag_bits = []
        for f in r["flags"]:
            flag_bits.append(
                f'<font color="{q["checks"][f]["color"]}">&bull; {q["checks"][f]["label"]}</font>')
            if f == "duplicate" and r.get("duplicates"):
                links = ", ".join(
                    f'<a href="{d["url"]}" color="#0a5bb5">{_html.escape(d["key"])}</a>'
                    for d in r["duplicates"])
                flag_bits.append(f'<font color="#555555">(=</font> {links}<font color="#555555">)</font>')
        flags_html = " ".join(flag_bits) or '<font color="#27ae60">Conforme</font>'
        idle = r["days_idle"]
        idle_html = (f'<font color="{_idle_hex(idle)}"><b>{idle}</b></font>'
                     if idle >= 0 else "—")
        st_color = STATUS_COLORS.get(r["status_category"], "#8895a7")
        data.append([
            Paragraph(f'<font color="white"><b>{r["flag_count"]}</b></font>', cellc),
            Paragraph(f'<a href="{r["url"]}" color="#0a5bb5"><u>{_html.escape(r["key"])}</u></a>', cell),
            Paragraph(_html.escape((r["summary"] or "")[:85]), cell),
            Paragraph(_html.escape(r["assignee"] or "—"), cell),
            Paragraph(f'<font color="{st_color}">{_html.escape(r["status"] or "")}</font>', cell),
            Paragraph(idle_html, cellc),
            Paragraph(_people_pdf(r), cell),
            Paragraph(flags_html, cell),
        ])
        # Compteur : fond rouge si defaut, vert sinon
        fc = RED_HEX if r["flag_count"] else GREEN_HEX
        style_cmds.append(("BACKGROUND", (0, idx), (0, idx), colors.HexColor(fc)))
        if idx % 2 == 0:
            style_cmds.append(("BACKGROUND", (1, idx), (-1, idx), colors.HexColor("#f6f9fc")))

    tbl = Table(data, colWidths=[1.0 * cm, 1.9 * cm, 5.0 * cm, 2.6 * cm, 1.9 * cm,
                                 1.2 * cm, 4.8 * cm, 7.0 * cm],
                repeatRows=1)
    tbl.setStyle(TableStyle(style_cmds))
    elems.append(tbl)
    doc.build(elems)


# ----------------------------- XLSX -----------------------------
def _argb(hex_color: str) -> str:
    """#RRGGBB -> AARRGGBB (opaque) pour openpyxl."""
    return "FF" + hex_color.lstrip("#").upper()


def build_xlsx(q: dict, title: str, path: str) -> None:
    """Export Excel colore, fidele a l'affichage (couleurs des defauts, etats, inactivite)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    checks = q["checks"]
    blue = _argb(BLUE_HEX)
    white = "FFFFFFFF"
    header_font = Font(bold=True, color=white, size=11)
    header_fill = PatternFill("solid", fgColor=blue)
    top = Alignment(vertical="top", wrap_text=True)
    center = Alignment(horizontal="center", vertical="center")

    wb = Workbook()

    # ---- Feuille 1 : Synthese ----
    ws0 = wb.active
    ws0.title = "Synthèse"
    ws0["A1"] = "Qualité & hygiène des tickets"
    ws0["A1"].font = Font(bold=True, size=16, color=blue)
    ws0["A2"] = (f"{title} — {datetime.now().strftime('%d/%m/%Y %H:%M')} — "
                 f"{q['total']} tickets analysés, {q['with_defects']} avec défaut(s), "
                 f"{q['clean']} conformes — seuil inactivité {q['stale_days']} j")
    ws0["A2"].font = Font(size=10, color="FF555555")
    hc = ws0.cell(4, 1, "Contrôle"); hc.font = header_font; hc.fill = header_fill
    hn = ws0.cell(4, 2, "Nb tickets"); hn.font = header_font; hn.fill = header_fill
    for i, (code, c) in enumerate(checks.items(), start=5):
        ws0.cell(i, 1, c["label"])
        cell = ws0.cell(i, 2, c["count"])
        cell.font = Font(bold=True, color=_argb(c["color"]))
        cell.alignment = center
    ws0.column_dimensions["A"].width = 34
    ws0.column_dimensions["B"].width = 12
    ws0.freeze_panes = "A5"

    # ---- Feuille 2 : Tickets ----
    ws = wb.create_sheet("Tickets")
    headers = ["Déf.", "Ticket", "Sujet", "Assigné", "État", "Inactif (j)",
               "Arborescence", "Personnes", "Contrôles échoués", "Doublons"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    zebra = PatternFill("solid", fgColor="FFF6F9FC")
    for i, r in enumerate(q["issues"], start=2):
        # Def. : compteur, fond rouge/vert
        fc = ws.cell(i, 1, r["flag_count"])
        fc.fill = PatternFill("solid",
                              fgColor=_argb(RED_HEX if r["flag_count"] else GREEN_HEX))
        fc.font = Font(bold=True, color=white)
        fc.alignment = center
        # Ticket : lien hypertexte
        kc = ws.cell(i, 2, r["key"])
        if r.get("url"):
            kc.hyperlink = r["url"]
        kc.font = Font(color="FF0A5BB5", bold=True, underline="single")
        # Sujet
        ws.cell(i, 3, r["summary"] or "").alignment = top
        # Assigne (personne assignee)
        ac = ws.cell(i, 4, r["assignee"] or "—")
        ac.alignment = top
        # Etat : couleur de categorie
        st = ws.cell(i, 5, r["status"] or "")
        st.font = Font(bold=True,
                       color=_argb(STATUS_COLORS.get(r["status_category"], "#8895a7")))
        # Inactivite : couleur selon anciennete
        idle = r["days_idle"]
        ic = ws.cell(i, 6, idle if idle >= 0 else None)
        if idle >= 0:
            ic.font = Font(color=_argb(_idle_hex(idle)), bold=idle >= 30)
        ic.alignment = center
        # Arborescence
        ws.cell(i, 7, _tree_text(r)).alignment = top
        # Personnes
        ws.cell(i, 8, _people_text(r)).alignment = top
        # Controles echoues : texte riche colore (une couleur par defaut)
        cell = ws.cell(i, 9)
        if r["flags"]:
            blocks = []
            for j, f in enumerate(r["flags"]):
                if j:
                    blocks.append(TextBlock(InlineFont(color="FF888888"), "  •  "))
                if f == "stale" and r["days_idle"] >= 0:
                    label = f"Non mis à jour depuis {r['days_idle']} j"
                elif f == "overdue" and r.get("days_overdue", -1) > 0:
                    label = f"Échéance dépassée depuis {r['days_overdue']} j"
                else:
                    label = checks[f]["label"]
                blocks.append(TextBlock(InlineFont(color=_argb(checks[f]["color"]),
                                                   b=True), label))
            cell.value = CellRichText(blocks)
        else:
            cell.value = "Conforme"
            cell.font = Font(color=_argb(GREEN_HEX), bold=True)
        cell.alignment = top
        # Doublons : lien cliquable si unique, sinon liste des cles
        dups = r.get("duplicates", []) or []
        dc = ws.cell(i, 10)
        if len(dups) == 1:
            dc.value = dups[0]["key"]
            if dups[0].get("url"):
                dc.hyperlink = dups[0]["url"]
            dc.font = Font(color="FF0A5BB5", bold=True, underline="single")
        elif dups:
            dc.value = ", ".join(d["key"] for d in dups)
        dc.alignment = top
        # Zebra (hors colonne 1 pour garder le rouge/vert)
        if i % 2 == 0:
            for c in range(2, 11):
                ws.cell(i, c).fill = zebra

    widths = [6, 13, 48, 22, 15, 11, 20, 42, 55, 22]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{len(q['issues']) + 1}"

    wb.save(path)
