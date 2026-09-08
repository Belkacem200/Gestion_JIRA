"""Generateur de trame PowerPoint (python-pptx) pour presentation hierarchie."""
from __future__ import annotations

from datetime import datetime

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .analysis import Analysis

BLUE = RGBColor(0x00, 0x3B, 0x7A)
RED = RGBColor(0xC0, 0x39, 0x2B)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
GREEN = RGBColor(0x27, 0xAE, 0x60)
GREY = RGBColor(0x55, 0x55, 0x55)

_SEV_COLOR = {"eleve": RED, "moyen": ORANGE, "faible": RGBColor(0xC9, 0xA2, 0x27)}


def _title_slide(prs: Presentation, title: str, author: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(9), Inches(2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = BLUE
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = datetime.now().strftime("%d/%m/%Y") + (f"  —  {author}" if author else "")
    r2.font.size = Pt(18); r2.font.color.rgb = GREY


def _section_title(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    p = box.text_frame.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = BLUE


def _bullets_slide(prs: Presentation, title: str, bullets: list[str],
                   colors: list[RGBColor] | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _section_title(slide, title)
    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(8.8), Inches(5.5))
    tf = box.text_frame; tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = f"•  {b}"
        r.font.size = Pt(16)
        if colors and i < len(colors):
            r.font.color.rgb = colors[i]
        p.space_after = Pt(8)


def _kpi_slide(prs: Presentation, a: Analysis) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _section_title(slide, "Avancement global")
    kpis = [
        (f"{a.progress_pct}%", "Avancement", GREEN if a.progress_pct >= 60 else ORANGE),
        (str(a.total), "Tickets", BLUE),
        (str(a.done), "Terminés", GREEN),
        (str(a.in_progress), "En cours", ORANGE),
        (str(a.todo), "À faire", GREY),
    ]
    x = 0.5
    for value, label, color in kpis:
        card = slide.shapes.add_textbox(Inches(x), Inches(2.0), Inches(1.75), Inches(1.6))
        tf = card.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = value
        r.font.size = Pt(36); r.font.bold = True; r.font.color.rgb = color
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = label
        r2.font.size = Pt(13); r2.font.color.rgb = GREY
        x += 1.85
    note = slide.shapes.add_textbox(Inches(0.6), Inches(4.2), Inches(9), Inches(1))
    pr = note.text_frame.paragraphs[0].add_run()
    pr.text = (f"Dynamique {a.momentum} — {a.resolved_in_period} résolus / "
               f"{a.created_in_period} créés sur {a.lookback_days} j")
    pr.font.size = Pt(15); pr.font.italic = True; pr.font.color.rgb = GREY


def _table_slide(prs: Presentation, title: str, headers: list[str],
                 rows: list[list[str]], col_widths: list[float] | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _section_title(slide, title)
    n_rows = len(rows) + 1
    n_cols = len(headers)
    tbl_shape = slide.shapes.add_table(
        n_rows, n_cols, Inches(0.4), Inches(1.3), Inches(9.2), Inches(0.4 * n_rows))
    table = tbl_shape.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        para = cell.text_frame.paragraphs[0]
        para.runs[0].font.size = Pt(12); para.runs[0].font.bold = True
        para.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(val)
            para = cell.text_frame.paragraphs[0]
            if para.runs:
                para.runs[0].font.size = Pt(11)


def build(a: Analysis, title: str, author: str, path: str) -> None:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    _title_slide(prs, title, author)
    _bullets_slide(prs, "Synthèse exécutive", a.highlights)
    _kpi_slide(prs, a)

    # Themes
    theme_rows = [
        [t.name[:28], f"{t.progress_pct}%", str(t.done), str(t.in_progress),
         str(t.todo), str(t.blockers)]
        for t in a.themes[:10]
    ]
    _table_slide(prs, "Sujets en cours (par thème)",
                 ["Thème", "%", "Fait", "En cours", "À faire", "Bloq."],
                 theme_rows, [3.2, 1.0, 1.0, 1.3, 1.3, 1.0])

    # Roadmap - En cours
    en_cours = [r for r in a.roadmap if r.phase == "En cours"]
    if en_cours:
        rows = [
            [r.name[:30], f"{r.progress_pct}%", str(r.in_progress), str(r.todo),
             r.next_due.strftime("%d/%m/%y") if r.next_due else "—",
             str(r.blockers) if r.blockers else "-"]
            for r in en_cours[:10]
        ]
        _table_slide(prs, "Roadmap — Chantiers en cours",
                     ["Chantier", "%", "En cours", "Reste", "Échéance", "Bloq."],
                     rows, [3.4, 0.9, 1.2, 1.0, 1.4, 0.9])

    # Roadmap - A venir
    a_venir = [r for r in a.roadmap if r.phase == "A venir"]
    if a_venir:
        rows = [
            [r.name[:40], str(r.todo),
             r.next_due.strftime("%d/%m/%y") if r.next_due else "—"]
            for r in a_venir[:12]
        ]
        _table_slide(prs, "Roadmap — Sujets à venir",
                     ["Chantier", "À faire", "Échéance"],
                     rows, [5.8, 1.6, 1.6])

    # Risques
    if a.risks:
        risk_rows = [
            [r.severity, r.key, (r.summary[:32] + "…") if len(r.summary) > 32 else r.summary,
             (r.reason[:40] + "…") if len(r.reason) > 40 else r.reason]
            for r in a.risks[:10]
        ]
        _table_slide(prs, "Points bloquants & risques",
                     ["Sév.", "Ticket", "Sujet", "Motif"],
                     risk_rows, [1.2, 1.4, 3.2, 3.4])

    # Recommandations
    from .report_markdown import _recommendations
    _bullets_slide(prs, "Décisions / arbitrages demandés", _recommendations(a),
                   [RED] * 10)

    prs.save(path)
