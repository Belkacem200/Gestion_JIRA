"""Export de la roadmap (frise Gantt) en image PNG, PDF, PPTX, HTML et email.

Strategie : une image PNG maitre est dessinee avec Pillow (rendu fidele de la
frise), puis reutilisee pour le PDF, le PPTX et l'email. L'export HTML est un
rendu natif autonome (vectoriel, ouvrable seul).
"""
from __future__ import annotations

import html as _html
from datetime import date, datetime
from email.mime.image import MIMEImage
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from io import BytesIO

# --- Couleurs (identiques a la page web) ---
BLUE = "#003B7A"
BLUE2 = "#0a5bb5"
TODO = "#5b7fb0"
INPROG = "#e67e22"
DONE = "#27ae60"
OVER = "#c0392b"
RED = "#c0392b"
LINE = "#e2e8f0"
EPIC_TRACK = "#d7e2f2"
GREY = "#66748a"


# ----------------------------------------------------------------------------
# Helpers communs
# ----------------------------------------------------------------------------
def _d(iso: str) -> date:
    return date.fromisoformat(iso[:10])


def _fmt(iso: str | None) -> str:
    if not iso:
        return "—"
    try:
        return _d(iso).strftime("%d/%m/%Y")
    except ValueError:
        return "—"


def _status_hex(cat: str, overdue: bool) -> str:
    if overdue:
        return OVER
    return {"Done": DONE, "In Progress": INPROG}.get(cat, TODO)


def _status_label(cat: str) -> str:
    return {"Done": "Terminé", "In Progress": "En cours"}.get(cat, "À faire")


def _positioner(data: dict):
    """Retourne une fonction pos(iso)->fraction [0,1] calee sur les colonnes de mois."""
    months = data["months"]
    starts = [_d(m["iso"]) for m in months]
    win_end = _d(data["window"]["end"])
    n = len(starts)

    def pos(iso: str | None) -> float:
        if not iso or n == 0:
            return 0.0
        d = _d(iso)
        if d <= starts[0]:
            return 0.0
        for i in range(n):
            s = starts[i]
            e = starts[i + 1] if i + 1 < n else win_end
            if d < e:
                return (i + (d - s).days / max(1, (e - s).days)) / n
        return 1.0

    return pos, n


def _rows(data: dict) -> list[tuple]:
    """Liste plate des lignes a dessiner : ('group'|'epic'|'item', payload)."""
    rows: list[tuple] = []
    lanes = data.get("lanes", [])
    solo = data.get("standalone", [])
    rows.append(("group", f"EPICS ({len(lanes)})"))
    for e in lanes:
        rows.append(("epic", e))
    if solo:
        rows.append(("group", f"HORS EPIC — PLANIFIES ({len(solo)})"))
        for s in solo:
            rows.append(("item", s))
    return rows


def _meta_text(n: dict, is_epic: bool) -> str:
    if is_epic:
        return f"{n.get('child_done', 0)}/{n.get('child_total', 0)} terminés · échéance {_fmt(n.get('due'))}"
    who = f" · {n['assignee']}" if n.get("assignee") else ""
    return f"{_status_label(n['status_category'])}{who} · échéance {_fmt(n.get('due'))}"


# ----------------------------------------------------------------------------
# Rendu PNG (Pillow) — image maitre
# ----------------------------------------------------------------------------
def _font(size: int, bold: bool = False):
    from PIL import ImageFont
    names = (["arialbd.ttf", "seguisb.ttf", "segoeuib.ttf"] if bold
             else ["arial.ttf", "segoeui.ttf"])
    for nm in names:
        try:
            return ImageFont.truetype(nm, size)
        except OSError:
            continue
    return ImageFont.load_default()


def _trunc(draw, text: str, font, maxw: float) -> str:
    text = text or ""
    if draw.textlength(text, font=font) <= maxw:
        return text
    while text and draw.textlength(text + "…", font=font) > maxw:
        text = text[:-1]
    return (text + "…") if text else ""


def _render_png_bytes(data: dict, title: str, scale: int = 2) -> bytes:
    from PIL import Image, ImageDraw

    S = scale
    LW, MW, RH = 300 * S, 96 * S, 34 * S
    HDR, TITLE_H, FOOT, PADX = 46 * S, 100 * S, 34 * S, 24 * S
    months = data.get("months", [])
    rows = _rows(data)
    pos, n = _positioner(data)

    chart_w = LW + MW * max(n, 1)
    W = PADX * 2 + chart_w
    H = TITLE_H + HDR + RH * max(len(rows), 1) + FOOT

    img = Image.new("RGB", (W, H), "#ffffff")
    d = ImageDraw.Draw(img)

    f_title = _font(22 * S, bold=True)
    f_sub = _font(12 * S)
    f_mo = _font(12 * S, bold=True)
    f_key = _font(12 * S, bold=True)
    f_sm = _font(11 * S)
    f_meta = _font(10 * S)
    f_cap = _font(11 * S, bold=True)

    x0 = PADX
    c = data.get("counts", {})

    # Titre + sous-titre
    d.text((x0, 16 * S), f"{title} — Roadmap", fill=BLUE, font=f_title, anchor="lm")
    sub = (f"Généré le {data.get('generated_at', '')} · {c.get('epics', 0)} épics · "
           f"{c.get('in_progress', 0)} en cours · {c.get('done', 0)} terminés · "
           f"{c.get('overdue', 0)} en retard · {c.get('no_date', 0)} sans échéance")
    d.text((x0, 46 * S), sub, fill=GREY, font=f_sub, anchor="lm")

    # Legende
    lx, ly = x0, 74 * S
    for label, color in [("À faire", TODO), ("En cours", INPROG),
                         ("Terminé", DONE), ("En retard", OVER)]:
        d.rounded_rectangle([lx, ly - 6 * S, lx + 16 * S, ly + 5 * S],
                            radius=2 * S, fill=color)
        d.text((lx + 20 * S, ly), label, fill=GREY, font=f_meta, anchor="lm")
        lx += 20 * S + d.textlength(label, font=f_meta) + 22 * S
    d.line([(lx, ly - 7 * S), (lx, ly + 6 * S)], fill=RED, width=2 * S)
    d.text((lx + 6 * S, ly), "aujourd'hui", fill=RED, font=f_meta, anchor="lm")

    top = TITLE_H
    chart_x = x0 + LW
    bottom = H - FOOT

    # Colonnes de mois + en-tete
    d.rectangle([x0, top, chart_x, top + HDR], fill="#eef3fa")
    d.text((x0 + 8 * S, top + HDR / 2), "Épic / Ticket", fill=BLUE, font=f_mo, anchor="lm")
    for i, m in enumerate(months):
        mx = chart_x + i * MW
        d.line([(mx, top), (mx, bottom)], fill=LINE, width=max(1, S // 2))
        d.rectangle([mx, top, mx + MW, top + HDR], fill="#eef3fa")
        d.text((mx + MW / 2, top + HDR / 2), m["label"], fill=BLUE, font=f_mo, anchor="mm")
    d.line([(chart_x + MW * n, top), (chart_x + MW * n, bottom)], fill=LINE, width=max(1, S // 2))

    # Ligne "aujourd'hui"
    tx = chart_x + pos(data.get("today")) * MW * n
    d.line([(tx, top + HDR), (tx, bottom)], fill=RED, width=2 * S)

    # Lignes
    ry = top + HDR
    if not [r for r in rows if r[0] != "group"]:
        d.text((x0 + 8 * S, ry + RH), "Aucune donnée à afficher.", fill=GREY, font=f_sm, anchor="lm")

    for kind, payload in rows:
        if kind == "group":
            d.rectangle([x0, ry, x0 + chart_w, ry + RH], fill="#f4f7fb")
            d.text((x0 + 8 * S, ry + RH / 2), payload, fill="#556", font=f_key, anchor="lm")
            ry += RH
            continue

        node = payload
        is_epic = (kind == "epic")
        # Libelle
        lx = x0 + 8 * S
        d.text((lx, ry + RH * 0.34), node["key"], fill=BLUE, font=f_key, anchor="lm")
        kw = d.textlength(node["key"] + "  ", font=f_key)
        sm = _trunc(d, node.get("summary", ""), f_sm, LW - 18 * S - kw)
        d.text((lx + kw, ry + RH * 0.34), sm, fill="#33415a", font=f_sm, anchor="lm")
        meta = _trunc(d, _meta_text(node, is_epic), f_meta, LW - 18 * S)
        d.text((lx, ry + RH * 0.72), meta, fill="#78849a", font=f_meta, anchor="lm")

        # Barre
        bx = chart_x + pos(node["start"]) * MW * n
        bx2 = chart_x + pos(node["end"]) * MW * n
        bw = max(4 * S, bx2 - bx)
        overdue = bool(node.get("overdue"))
        col = _status_hex(node["status_category"], overdue)
        if is_epic:
            bh = 20 * S
            by = ry + (RH - bh) / 2
            d.rounded_rectangle([bx, by, bx + bw, by + bh], radius=5 * S, fill=EPIC_TRACK)
            prog = node.get("progress") or 0
            if prog > 0:
                pw = max(4 * S, bw * prog)
                d.rounded_rectangle([bx, by, bx + pw, by + bh], radius=5 * S, fill=DONE)
            cap = f"{node['key']} · {round(prog * 100)}%"
            d.text((bx + 7 * S, by + bh / 2), _trunc(d, cap, f_cap, bw - 10 * S),
                   fill="#12315e", font=f_cap, anchor="lm")
        else:
            bh = 15 * S
            by = ry + (RH - bh) / 2
            d.rounded_rectangle([bx, by, bx + bw, by + bh], radius=4 * S, fill=col)
            if overdue:
                d.rounded_rectangle([bx, by, bx + bw, by + bh], radius=4 * S,
                                    outline=RED, width=2 * S)
            if bw > 56 * S:
                d.text((bx + 6 * S, by + bh / 2), _trunc(d, node["key"], f_cap, bw - 10 * S),
                       fill="#ffffff", font=f_cap, anchor="lm")
            else:
                d.text((bx + bw + 5 * S, by + bh / 2), node["key"],
                       fill="#33415a", font=f_meta, anchor="lm")

        d.line([(x0, ry + RH), (x0 + chart_w, ry + RH)], fill="#eef1f6", width=max(1, S // 2))
        ry += RH

    d.text((x0, H - FOOT + 12 * S),
           "Roadmap BACNSO — générée automatiquement depuis JIRA (création → échéance)",
           fill="#9aa4b2", font=f_meta, anchor="lm")

    buf = BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def build_png(data: dict, title: str, path: str) -> str:
    with open(path, "wb") as fh:
        fh.write(_render_png_bytes(data, title, scale=2))
    return path


# ----------------------------------------------------------------------------
# PDF (image maitre placee sur une page a sa dimension)
# ----------------------------------------------------------------------------
def build_pdf(data: dict, title: str, path: str) -> str:
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader

    png = _render_png_bytes(data, title, scale=2)
    reader = ImageReader(BytesIO(png))
    iw, ih = reader.getSize()
    target_w = 1400.0                      # largeur de page en points
    sc = target_w / iw
    page = (target_w, ih * sc)
    c = canvas.Canvas(path, pagesize=page)
    c.drawImage(reader, 0, 0, width=page[0], height=page[1])
    c.showPage()
    c.save()
    return path


# ----------------------------------------------------------------------------
# PPTX (image maitre inseree sur une diapo)
# ----------------------------------------------------------------------------
def build_pptx(data: dict, title: str, path: str) -> str:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    from PIL import Image

    png = _render_png_bytes(data, title, scale=2)
    iw, ih = Image.open(BytesIO(png)).size

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])   # vierge

    tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = f"{title} — Roadmap"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x3B, 0x7A)

    max_w = Inches(13.0)
    max_h = Inches(6.7)
    top = Inches(0.75)
    sc = min(max_w / iw, max_h / ih)
    w = int(iw * sc)
    h = int(ih * sc)
    left = int((prs.slide_width - w) / 2)
    slide.shapes.add_picture(BytesIO(png), left, top, width=w, height=h)
    prs.save(path)
    return path


# ----------------------------------------------------------------------------
# HTML natif autonome (frise vectorielle)
# ----------------------------------------------------------------------------
def _bar_html(n: dict, is_epic: bool, pos) -> str:
    left = max(0.0, pos(n["start"]) * 100)
    width = max(0.5, (pos(n["end"]) - pos(n["start"])) * 100)
    over = " over" if n.get("overdue") else ""
    color = _status_hex(n["status_category"], n.get("overdue"))
    url = _html.escape(n.get("url", ""))
    if is_epic:
        prog = round((n.get("progress") or 0) * 100)
        inner = (f'<i style="width:{prog}%"></i>'
                 f'<span class="cap">{_html.escape(n["key"])} · {prog}%</span>')
        return (f'<a class="bar epic{over}" href="{url}" target="_blank" '
                f'style="left:{left:.3f}%;width:{width:.3f}%">{inner}</a>')
    cap = (f'<span class="cap">{_html.escape(n["key"])}</span>' if width > 8 else "")
    out = (f'<span class="out" style="left:{min(left + width, 99):.3f}%">'
           f'{_html.escape(n["key"])}</span>' if width <= 8 else "")
    return (f'<a class="bar{over}" href="{url}" target="_blank" '
            f'style="left:{left:.3f}%;width:{width:.3f}%;--c:{color}">{cap}</a>{out}')


def _label_html(n: dict, is_epic: bool) -> str:
    url = _html.escape(n.get("url", ""))
    return (f'<div class="label"><a href="{url}" target="_blank">{_html.escape(n["key"])}</a>'
            f'<span class="s" title="{_html.escape(n.get("summary", ""))}">'
            f'{_html.escape(n.get("summary", ""))}</span>'
            f'<span class="m">{_html.escape(_meta_text(n, is_epic))}</span></div>')


def build_html(data: dict, title: str) -> str:
    pos, n = _positioner(data)
    months = data.get("months", [])
    lanes = data.get("lanes", [])
    solo = data.get("standalone", [])
    c = data.get("counts", {})

    months_html = "".join(
        f'<div class="mo">{_html.escape(m["label"])}</div>' for m in months)

    body = [f'<div class="grp">EPICS ({len(lanes)})</div>']
    for e in lanes:
        body.append(
            f'<div class="row">{_label_html(e, True)}'
            f'<div class="track">{_bar_html(e, True, pos)}</div></div>')
        kids = e.get("children", [])
        if kids:
            krows = "".join(
                f'<div class="krow">{_label_html(k, False)}'
                f'<div class="track">{_bar_html(k, False, pos)}</div></div>'
                for k in kids)
            body.append(
                f'<details class="kids"><summary>▸ {len(kids)} ticket(s) enfant(s)</summary>'
                f'{krows}</details>')
    if solo:
        body.append(f'<div class="grp">HORS ÉPIC — PLANIFIÉS ({len(solo)})</div>')
        for s in solo:
            body.append(
                f'<div class="row">{_label_html(s, False)}'
                f'<div class="track">{_bar_html(s, False, pos)}</div></div>')

    today_left = pos(data.get("today")) * 100
    stamp = data.get("generated_at", "")
    return f"""<!DOCTYPE html><html lang="fr"><head><meta charset="UTF-8"/>
<title>Roadmap — {_html.escape(title)}</title><style>
:root{{--blue:#003B7A;--blue2:#0a5bb5;--line:#e2e8f0;--todo:#5b7fb0;--orange:#e67e22;
--green:#27ae60;--red:#c0392b;--lw:300px;--mw:94px;--cols:{max(n,1)};}}
*{{box-sizing:border-box;}}body{{font-family:'Segoe UI',Arial,sans-serif;margin:0;
background:#f4f7fb;color:#1a2233;}}.wrap{{padding:18px 22px;}}h1{{color:var(--blue);
font-size:20px;margin:0 0 2px;}}.sub{{color:#667;font-size:12px;}}
.legend{{display:flex;gap:16px;font-size:12px;color:#667;margin:10px 0 14px;flex-wrap:wrap;}}
.legend span{{display:inline-flex;align-items:center;gap:5px;}}
.lg{{width:14px;height:10px;border-radius:3px;display:inline-block;}}
.lgt{{width:2px;height:14px;background:var(--red);display:inline-block;}}
.chart{{background:#fff;border-radius:10px;box-shadow:0 1px 4px rgba(0,0,0,.08);overflow-x:auto;}}
.inner{{width:calc(var(--lw) + var(--mw)*var(--cols));min-width:100%;position:relative;}}
.head{{display:flex;background:#eef3fa;border-bottom:1px solid var(--line);position:sticky;top:0;z-index:3;}}
.corner{{flex:0 0 var(--lw);width:var(--lw);padding:9px 12px;font-weight:700;color:var(--blue);
font-size:12px;border-right:1px solid var(--line);}}
.months{{display:flex;flex:1;}}.mo{{flex:1;min-width:var(--mw);text-align:center;padding:9px 2px;
font-size:11.5px;color:var(--blue);border-left:1px solid var(--line);white-space:nowrap;}}
.bodyc{{position:relative;}}
.overlay{{position:absolute;top:0;bottom:0;left:var(--lw);right:0;pointer-events:none;}}
.today{{position:absolute;top:0;bottom:0;width:2px;background:var(--red);left:{today_left:.3f}%;}}
.today::after{{content:'aujourd’hui';position:absolute;top:1px;left:3px;font-size:10px;
color:var(--red);font-weight:700;background:rgba(255,255,255,.85);padding:0 3px;border-radius:3px;white-space:nowrap;}}
.grp{{padding:10px 12px 4px;font-size:12px;font-weight:700;color:#667;text-transform:uppercase;
letter-spacing:.4px;background:#f4f7fb;}}
.row,.krow{{display:flex;align-items:stretch;border-bottom:1px solid #eef1f6;min-height:38px;}}
.krow{{min-height:30px;background:#fbfcfe;}}
.label{{flex:0 0 var(--lw);width:var(--lw);padding:6px 12px;border-right:1px solid var(--line);overflow:hidden;}}
.krow .label{{padding-left:26px;}}
.label a{{color:var(--blue2);font-weight:700;text-decoration:none;font-size:12.5px;white-space:nowrap;}}
.label .s{{font-size:12px;color:#33415a;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;display:block;}}
.label .m{{font-size:10.5px;color:#78849a;margin-top:1px;}}
.track{{position:relative;flex:1;min-width:calc(var(--mw)*var(--cols));
background-image:linear-gradient(90deg,var(--line) 1px,transparent 1px);background-size:var(--mw) 100%;}}
.bar{{position:absolute;top:50%;transform:translateY(-50%);height:16px;border-radius:5px;
background:var(--c,#5b7fb0);display:flex;align-items:center;text-decoration:none;
box-shadow:0 1px 3px rgba(0,0,0,.18);overflow:hidden;min-width:4px;}}
.bar.epic{{height:22px;border-radius:7px;background:#d7e2f2;}}
.bar.epic>i{{display:block;height:100%;background:linear-gradient(90deg,var(--blue2),var(--green));border-radius:inherit;}}
.bar .cap{{position:absolute;left:8px;font-size:10.5px;font-weight:700;color:#fff;
white-space:nowrap;text-shadow:0 1px 1px rgba(0,0,0,.35);}}
.bar.epic .cap{{color:#12315e;text-shadow:none;}}
.bar.over{{outline:2px solid var(--red);outline-offset:-2px;}}
.out{{position:absolute;top:50%;transform:translateY(-50%);font-size:10.5px;color:#33415a;white-space:nowrap;padding-left:5px;}}
details.kids>summary{{list-style:none;cursor:pointer;padding:4px 12px 4px 26px;font-size:11.5px;
color:var(--blue2);font-weight:600;background:#f7faff;border-bottom:1px solid #eef1f6;}}
details.kids>summary::-webkit-details-marker{{display:none;}}
.foot{{color:#9aa4b2;font-size:11px;margin-top:10px;}}
</style></head><body><div class="wrap">
<h1>🗺️ {_html.escape(title)} — Roadmap</h1>
<div class="sub">Générée le {_html.escape(stamp)} · {c.get('epics',0)} épics · {c.get('in_progress',0)} en cours · {c.get('done',0)} terminés · {c.get('overdue',0)} en retard · {c.get('no_date',0)} sans échéance</div>
<div class="legend">
<span><i class="lg" style="background:var(--todo)"></i> À faire</span>
<span><i class="lg" style="background:var(--orange)"></i> En cours</span>
<span><i class="lg" style="background:var(--green)"></i> Terminé</span>
<span><i class="lg" style="background:var(--red)"></i> En retard</span>
<span><i class="lgt"></i> aujourd'hui</span></div>
<div class="chart"><div class="inner">
<div class="head"><div class="corner">Épic / Ticket</div><div class="months">{months_html}</div></div>
<div class="bodyc"><div class="overlay"><div class="today"></div></div>
{''.join(body)}
</div></div></div>
<p class="foot">Chaque barre va de la création à l'échéance (ou à la résolution / aujourd'hui si aucune échéance). Généré automatiquement depuis JIRA.</p>
</div></body></html>"""


# ----------------------------------------------------------------------------
# Email (.eml) : image de la roadmap embarquee (cid) + synthese
# ----------------------------------------------------------------------------
def _email_html(data: dict, title: str) -> str:
    c = data.get("counts", {})
    lanes = data.get("lanes", [])
    # Prochaines echeances (epics non termines a 100%, avec echeance, triees)
    upcoming = sorted(
        [e for e in lanes if e.get("due") and (e.get("progress") or 0) < 1.0],
        key=lambda e: e["due"])[:8]
    rows = "".join(
        f"<tr><td style='padding:4px 8px'><a href='{_html.escape(e.get('url',''))}'>"
        f"{_html.escape(e['key'])}</a></td>"
        f"<td style='padding:4px 8px'>{_html.escape(e.get('summary',''))}</td>"
        f"<td style='padding:4px 8px;color:"
        f"{OVER if e.get('overdue') else '#333'}'>{_fmt(e.get('due'))}</td>"
        f"<td style='padding:4px 8px'>{round((e.get('progress') or 0)*100)}%</td></tr>"
        for e in upcoming) or "<tr><td colspan='4' style='padding:4px 8px'>—</td></tr>"
    return f"""\
<html><body style="font-family:Segoe UI,Arial,sans-serif;color:#222;line-height:1.5">
  <h2 style="color:#003B7A;margin-bottom:2px">{_html.escape(title)} — Roadmap</h2>
  <p style="color:#666;margin-top:0">Générée le {_html.escape(data.get('generated_at',''))}</p>
  <div style="background:#F4F7FB;border-left:4px solid #003B7A;padding:10px 14px;margin:12px 0">
    <b>{c.get('epics',0)} épics</b> — {c.get('in_progress',0)} en cours,
    {c.get('done',0)} terminés,
    <span style="color:#C0392B">{c.get('overdue',0)} en retard</span>,
    {c.get('no_date',0)} sans échéance.
  </div>
  <img src="cid:roadmap" alt="Roadmap" style="max-width:100%;border:1px solid #e2e8f0;border-radius:8px"/>
  <h3 style="color:#003B7A">Prochaines échéances</h3>
  <table style="border-collapse:collapse;font-size:14px">
    <tr style="background:#003B7A;color:#fff">
      <th style="padding:4px 8px;text-align:left">Épic</th>
      <th style="padding:4px 8px;text-align:left">Sujet</th>
      <th style="padding:4px 8px;text-align:left">Échéance</th>
      <th style="padding:4px 8px;text-align:left">Avanct</th>
    </tr>{rows}
  </table>
  <p style="color:#888;font-size:12px;margin-top:20px">
    Roadmap générée automatiquement depuis JIRA.
  </p>
</body></html>"""


def build_eml(data: dict, title: str, path: str) -> str:
    png = _render_png_bytes(data, title, scale=2)
    html = _email_html(data, title)
    root = MIMEMultipart("related")
    root["Subject"] = f"{title} — Roadmap — {datetime.now().strftime('%d/%m/%Y')}"
    root["From"] = "moi"
    root["To"] = ""
    alt = MIMEMultipart("alternative")
    root.attach(alt)
    alt.attach(MIMEText("Roadmap BACNSO — voir la version HTML (image intégrée).",
                        "plain", "utf-8"))
    alt.attach(MIMEText(html, "html", "utf-8"))
    img = MIMEImage(png, _subtype="png")
    img.add_header("Content-ID", "<roadmap>")
    img.add_header("Content-Disposition", "inline", filename="roadmap.png")
    root.attach(img)
    with open(path, "wb") as fh:
        fh.write(root.as_bytes())
    return path
